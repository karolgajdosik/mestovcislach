"""Hospodárenie krajských miest – porovnateľné údaje pre všetky mestá.

Zdroje:
  1. Register účtovných závierok (registeruz.sk, otvorené API): individuálne účtovné závierky miest (nie konsolidované)
     – výkaz ziskov a strát (Úč ROPO SFOV 2-01) a súvaha (Úč ROPO SFOV 1-01): náklady, výnosy, daňové výnosy,
     transfery, výsledok hospodárenia, majetok, záväzky, bankové úvery.
  2. MF SR – Dlh a dlhová služba miest a obcí SR k ultimu roka (opendata.mfsr.sk): dlh podľa zákona, bežné príjmy,
     pomer dlhu (limit 60 %) a dlhovej služby (limit 25 %). Bratislava a Košice = magistrát bez mestských častí.

Výstup: data/mesta/hospodarenie.json
"""

from __future__ import annotations

import io
import json
import re
import time
from datetime import datetime, timezone

from spolocne import DATA, KOREN, s, spusti, zapis_ak_zmena
from zber_ukazovatele import najdi, stlpce

RUZ = "https://www.registeruz.sk/cruz-public/api/"
URL_DLH = "https://opendata.mfsr.sk/opendata/download/dlh-a-dlhova-sluzba-miest-a-obci-sr-k-ultimu-roka/{}"
MESTA = json.loads((KOREN / "src" / "mesta.json").read_text(encoding="utf-8"))
OD_ROKU = 2015
# Mestské časti Bratislavy (IČO overené v RÚZ a CRZ) – sekcia na /bratislava/financie
MESTSKE_CASTI = {
    "Staré Mesto": "00603147", "Ružinov": "00603155", "Vrakuňa": "00603295", "Podunajské Biskupice": "00641383", "Nové Mesto": "00603317",
    "Rača": "00304557", "Vajnory": "00304565", "Karlova Ves": "00603520", "Dúbravka": "00603406", "Lamač": "00603414", "Devín": "00603422",
    "Devínska Nová Ves": "00603392", "Záhorská Bystrica": "00604887", "Petržalka": "00603201", "Jarovce": "00304603", "Rusovce": "00304611",
    "Čunovo": "00641243",
}

# (kľúč, výkaz, tabuľka, číslo riadku, stĺpec) – stĺpce: VZaS [hlavná, podnikateľská, spolu, minulé obdobie],
# aktíva [brutto, korekcia, netto, netto minulé], pasíva [bežné, minulé]
UKAZOVATELE = [
    ("naklady", "zisk", "Náklady", 64, 2), ("osobne_naklady", "zisk", "Náklady", 11, 2), ("sluzby", "zisk", "Náklady", 6, 2),
    ("spotreba", "zisk", "Náklady", 1, 2), ("odpisy", "zisk", "Náklady", 30, 2), ("transfery_dane", "zisk", "Náklady", 54, 2), ("uroky", "zisk", "Náklady", 42, 2),
    ("vynosy", "zisk", "Výnosy", 134, 2), ("danove_vynosy", "zisk", "Výnosy", 81, 2), ("poplatky", "zisk", "Výnosy", 82, 2),
    ("trzby", "zisk", "Výnosy", 65, 2), ("transfery_prijate", "zisk", "Výnosy", 124, 2), ("vysledok", "zisk", "Výnosy", 138, 2),
    ("majetok", "suvaha", "Strana aktív", 1, 2), ("dlhodoby_hmotny", "suvaha", "Strana aktív", 11, 2), ("financne_ucty", "suvaha", "Strana aktív", 85, 2),
    ("vlastne_imanie", "suvaha", "Strana pasív", 116, 0), ("zavazky", "suvaha", "Strana pasív", 126, 0), ("bankove_uvery", "suvaha", "Strana pasív", 173, 0),
]
_sablony: dict[int, dict] = {}


def ruz(cesta: str, **params) -> dict:
    for pokus in range(6):
        try:
            r = s.get(RUZ + cesta, params=params, timeout=60)
            time.sleep(0.6)  # API odpája pri rýchlych dopytoch
            if r.status_code == 200:
                return r.json()
        except Exception:  # noqa: BLE001
            pass
        time.sleep(5 * (pokus + 1))
    raise RuntimeError(f"RÚZ {cesta} {params} nedostupné")


def sablona(sid: int) -> dict:
    """Pozícia riadku v tabuľke podľa čísla riadku a počet stĺpcov."""
    if sid not in _sablony:
        sab = ruz("sablona", id=sid)
        _sablony[sid] = {"nazov": sab.get("nazov", ""), "tabulky": {t["nazov"]["sk"]: {"stlpcov": t.get("pocetStlpcov"),
                         "riadky": {int(r["cisloRiadku"]): i for i, r in enumerate(t.get("riadky", []))}} for t in sab.get("tabulky", [])}}
    return _sablony[sid]


def zavierky(ico: str) -> dict:
    ids = ruz("uctovne-jednotky", **{"zmenene-od": "2000-01-01", "ico": ico}).get("id", [])
    if not ids:
        raise RuntimeError(f"RÚZ: účtovná jednotka {ico} nenájdená")
    uj = ruz("uctovna-jednotka", id=ids[0])
    vysl: dict[str, dict] = {}
    for zid in uj.get("idUctovnychZavierok", []):
        z = ruz("uctovna-zavierka", id=zid)
        rok = (z.get("obdobieDo") or "")[:4]
        if z.get("konsolidovana") or not (z.get("obdobieDo") or "").endswith("-12") or not rok.isdigit() or int(rok) < OD_ROKU:
            continue
        # pri opravnej závierke berieme najnovšiu
        if rok in vysl and (vysl[rok].get("datumPoslednejUpravy") or "") >= (z.get("datumPoslednejUpravy") or ""):
            continue
        vysl[rok] = z
    return vysl


def ukazovatele(z: dict) -> dict:
    hodnoty: dict[str, float] = {}
    for vid in z.get("idUctovnychVykazov", []):
        v = ruz("uctovny-vykaz", id=vid)
        tab = (v.get("obsah") or {}).get("tabulky")
        if not tab or not v.get("idSablony"):
            continue
        sab = sablona(v["idSablony"])
        druh = "zisk" if re.search(r"ziskov a strát", sab["nazov"], re.I) else "suvaha" if re.search(r"súvaha", sab["nazov"], re.I) else None
        if not druh or "ROPO" not in sab["nazov"]:
            continue
        data = {t["nazov"]["sk"]: t["data"] for t in tab}
        for kluc, d, tabulka, riadok, stlpec in UKAZOVATELE:
            st = sab["tabulky"].get(tabulka)
            if d != druh or not st or tabulka not in data or riadok not in st["riadky"]:
                continue
            n = len(data[tabulka]) // max(1, len(st["riadky"]))
            try:
                hodnoty[kluc] = round(float(data[tabulka][st["riadky"][riadok] * n + stlpec] or 0))
            except (IndexError, ValueError):
                pass
    return hodnoty


def dlh() -> dict[str, dict]:
    import openpyxl
    nazvy = {"bratislava": r"^bratislava( magistrát| hlavné mesto)?$", "kosice": r"^košice( magistrát)?$"}
    roky: dict[str, dict] = {}
    for n in range(1, 20):
        r = s.get(URL_DLH.format(n), timeout=120)
        if r.status_code != 200 or r.content[:2] != b"PK":
            continue
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        l60 = next((x for x in wb.sheetnames if "60" in x and "kraj" in x), None)
        l25 = next((x for x in wb.sheetnames if "25" in x and "kraj" in x), None)
        if not l60:
            continue
        rows60 = [list(x) for x in wb[l60].iter_rows(values_only=True)]
        rok = next((m.group(1) for r_ in rows60[:8] for v in r_ if isinstance(v, str) and (m := re.search(r"31\.12\.(\d{4})", v))), None)
        if not rok:
            continue
        data: dict[str, dict] = {}
        hl = stlpce(rows60, 3, 12)
        i_dlh, i_pomer = najdi(hl, r"celková suma"), najdi(hl, r"pomer")
        rows25 = [list(x) for x in wb[l25].iter_rows(values_only=True)] if l25 else []
        hl25 = stlpce(rows25, 3, 12) if rows25 else {}
        i_prijmy, i_sluzba = najdi(hl25, r"úhrn bež\w*\.? príjmov"), najdi(hl25, r"pomer dlhovej")
        for m in MESTA:
            vzor = nazvy.get(m["slug"], rf"^{re.escape(m['nazov'].lower())}$")
            r60 = next((x for x in rows60 if re.match(vzor, str(x[0] or "").strip().lower())), None)
            if not r60 or i_dlh is None:
                continue
            d = {"dlh": round(float(r60[i_dlh] or 0)), "pomer": round(float(r60[i_pomer] or 0), 2) if i_pomer is not None else None}
            r25 = next((x for x in rows25 if re.match(vzor, str(x[0] or "").strip().lower())), None)
            if r25:
                d["bezne_prijmy"] = round(float(r25[i_prijmy] or 0)) if i_prijmy is not None else None
                d["dlhova_sluzba"] = round(float(r25[i_sluzba] or 0), 2) if i_sluzba is not None else None
            data[m["slug"]] = d
        if len(data) >= 6:
            roky[rok] = data
            print(f"Dlh {rok}: " + ", ".join(f"{k} {v['dlh']:,}" for k, v in data.items()))
    return dict(sorted(roky.items()))


def roky_jednotky(ico: str) -> dict:
    roky = {}
    for rok, z in sorted(zavierky(ico).items()):
        u = ukazovatele(z)
        if u.get("naklady") and u.get("vynosy"):
            roky[rok] = u
    if not roky:
        raise RuntimeError("žiadna závierka s výkazmi")
    return roky


def main() -> int:
    stare_subor = DATA / "mesta" / "hospodarenie.json"
    stare = json.loads(stare_subor.read_text(encoding="utf-8")) if stare_subor.exists() else {}
    vystup = {"zdroje": {
        "ruz": {"nazov": "Register účtovných závierok – individuálne účtovné závierky miest", "url": "https://www.registeruz.sk"},
        "dlh": {"nazov": "MF SR – Dlh a dlhová služba miest a obcí SR k ultimu roka", "url": "https://opendata.mfsr.sk"}},
        "mesta": {}, "mestske_casti": {}, "dlh": {}}
    chyby = 0
    for m in MESTA:
        try:
            roky = roky_jednotky(m["ico"])
            vystup["mesta"][m["slug"]] = {"roky": roky}
            p = roky[max(roky)]
            print(f"{m['nazov']}: {len(roky)} rokov, {max(roky)}: náklady {p['naklady']:,}, výnosy {p['vynosy']:,}, úvery {p.get('bankove_uvery', 0):,}")
        except Exception as e:  # noqa: BLE001
            chyby += 1
            print(f"::warning title=Hospodárenie {m['nazov']}::{e}")
            if m["slug"] in stare.get("mesta", {}):
                vystup["mesta"][m["slug"]] = stare["mesta"][m["slug"]]
    chyby_mc = 0
    for mc, ico in MESTSKE_CASTI.items():
        try:
            vystup["mestske_casti"][mc] = {"ico": ico, "roky": roky_jednotky(ico)}
        except Exception as e:  # noqa: BLE001
            chyby_mc += 1
            print(f"::warning title=Hospodárenie MČ {mc}::{e}")
            if mc in stare.get("mestske_casti", {}):
                vystup["mestske_casti"][mc] = stare["mestske_casti"][mc]
    print(f"Mestské časti: {len(vystup['mestske_casti'])} (chyby {chyby_mc})")
    try:
        vystup["dlh"] = dlh()
    except Exception as e:  # noqa: BLE001
        chyby += 1
        print(f"::warning title=Dlh miest::{e}")
        vystup["dlh"] = stare.get("dlh", {})
    if chyby > 4 or not vystup["mesta"]:
        raise RuntimeError(f"Hospodárenie: {chyby} chýb")
    vystup["aktualizovane"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    zapis_ak_zmena(DATA / "mesta" / "hospodarenie.json", vystup)
    return 0


if __name__ == "__main__":
    spusti(main, "Hospodárenie miest")
