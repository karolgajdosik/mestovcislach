import requests, re, io, zipfile, subprocess, json, os
UA={"User-Agent":"mestovcislach.sk (github.com/karolgajdosik/mestovcislach)"}
for u in ["https://volbysr.sk/oso/sk/subory_na_stiahnutie.html","https://volbysr.sk/osk/sk/subory_na_stiahnutie.html"]:
    t=requests.get(u,headers=UA,timeout=40).text
    print("==",u); print("  ", sorted(set(re.findall(r'href="([^"]+\.(?:zip|xlsx|csv|json|xml))"',t))))
    print("  ", re.sub(r'\s+',' ',re.sub('<[^>]+>',' ',t))[1500:3500])
import openpyxl
for u in ["https://volbysr.sk/oso/files/OSO2026_SK_xlsx.zip","https://volbysr.sk/osk/files/OSK2026_SK_xlsx.zip"]:
    r=requests.get(u,headers=UA,timeout=60); print("==",u,r.status_code,len(r.content))
    if r.status_code!=200: continue
    z=zipfile.ZipFile(io.BytesIO(r.content))
    for n in z.namelist():
        try:
            wb=openpyxl.load_workbook(io.BytesIO(z.read(n)),read_only=True)
            ws=wb.worksheets[0]; rows=list(ws.iter_rows(values_only=True,max_row=3))
            print("  ",n, ws.max_row, [str(c)[:20] for c in (rows[0] if rows else [])][:14])
        except Exception as e: print("  ",n,"ERR",e)
os.makedirs("diag",exist_ok=True)
subprocess.run(["pip","install","-q","-r","scripts/requirements.txt"])
for cmd in [["python","scripts/zber_komunalne.py"],["python","scripts/zber_mesta.py","komunalne"]]:
    r=subprocess.run(cmd,capture_output=True,text=True); print("==",cmd, r.returncode); print(r.stdout[-2500:]); print(r.stderr[-2500:])
d=json.load(open("data/komunalne-bratislava.json")); print("BA volby:", [(v["id"], v.get("stav"), bool(v.get("oso")), v.get("spracovanie")) for v in d["volby"]])
for s in ["kosice","trencin"]:
    k=json.load(open(f"data/mesta/{s}.json"))["komunalne"]; print(s, {y:(v.get("spracovanie"), bool(v.get("primator"))) for y,v in k.items()})
print("HOTOVO")
